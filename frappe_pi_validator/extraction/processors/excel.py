from __future__ import annotations

import ast
import operator
import re
from typing import Any

import xlrd
from openpyxl import load_workbook

from frappe_pi_validator.extraction.models.extraction import DocumentResult, Sheet, Table
from frappe_pi_validator.extraction.services.table.generic import (
    GenericTableExtractor,
    table_value,
    raw_table,
    rows_to_table,
    table_score,
)
from frappe_pi_validator.extraction.services.table.models import OCRBlock
from frappe_pi_validator.extraction.services.table_service import TableService


FORMULA_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
FORMULA_SUM = re.compile(
    r"SUM\(\s*\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)",
    re.IGNORECASE,
)

_OPERATORS = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.Div: operator.truediv,
}


def _column_index(letters: str) -> int:
    index = 0
    for letter in letters.upper():
        index = index * 26 + ord(letter) - 64
    return index - 1


def _safe_arithmetic(expression: str) -> float | None:
    """Evaluate + - * / and brackets on numbers only."""

    def evaluate(node):
        if isinstance(node, ast.Expression):
            return evaluate(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.USub):
            return -evaluate(node.operand)
        if isinstance(node, ast.BinOp) and type(node.op) in _OPERATORS:
            return _OPERATORS[type(node.op)](
                evaluate(node.left), evaluate(node.right)
            )
        raise ValueError("unsupported formula")

    try:
        return float(evaluate(ast.parse(expression, mode="eval")))
    except (ValueError, SyntaxError, ZeroDivisionError, TypeError):
        return None


def evaluate_formulas(rows: list[list[Any]]) -> list[list[Any]]:
    """
    Formula cells without a value cached by Excel (files written by
    other programs) are evaluated from the sheet itself:

        =E6*F6, =B7+C7, =SUM(E6:E12)

    Formulas using other functions are left as text.
    """

    def cell(row: int, column: int, depth: int):
        if row < 0 or row >= len(rows) or column < 0 or column >= len(rows[row]):
            return None
        value = rows[row][column]
        if isinstance(value, str) and value.startswith("="):
            return formula(value, depth + 1)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return value
        return None

    def formula(text: str, depth: int):
        if depth > 20:
            return None

        expression = text[1:]

        def sum_range(match):
            first_col, first_row = _column_index(match.group(1)), int(match.group(2)) - 1
            last_col, last_row = _column_index(match.group(3)), int(match.group(4)) - 1
            total = 0.0
            for r in range(first_row, last_row + 1):
                for c in range(first_col, last_col + 1):
                    total += cell(r, c, depth) or 0
            return repr(total)

        expression = FORMULA_SUM.sub(sum_range, expression)

        missing = False

        def reference(match):
            nonlocal missing
            value = cell(int(match.group(2)) - 1, _column_index(match.group(1)), depth)
            if value is None:
                missing = True
                return "0"
            return repr(float(value))

        expression = FORMULA_REF.sub(reference, expression)

        if missing:
            return None

        return _safe_arithmetic(expression)

    result = []

    for row_index, row in enumerate(rows):
        new_row = []
        for column_index, value in enumerate(row):
            if isinstance(value, str) and value.startswith("="):
                evaluated = formula(value, 0)
                if evaluated is not None:
                    value = round(evaluated, 6)
                    if float(value).is_integer():
                        value = int(value)
            new_row.append(value)
        result.append(new_row)

    return result


class ExcelProcessor:
    HEADERS = [
        "item_no",
        "brand",
        "size",
        "pattern",
        "load_index",
        "speed_rating",
        "pr",
        "sidewall",
        "quantity",
        "unit",
        "unit_price",
        "amount",
        "currency",
        "qty_in_40hq",
        "container_ratio",
    ]

    def __init__(self):
        self.table_service = TableService()

    def supports(self, mime_type: str) -> bool:
        return mime_type in {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel.sheet.macroEnabled.12",
            "application/vnd.ms-excel",
        }

    # ---------------------------------------------------------
    # Header-mapped format: header text -> field
    #
    # Checked in order; the first matching field wins, so the
    # more specific keywords come first ("qty/40hq" before "qty").
    # ---------------------------------------------------------

    COLUMN_KEYWORDS = [
        ("qty_in_40hq", ("40hq",)),
        ("amount", ("total amount", "total price", "amount")),
        ("unit_price", ("unit price", "/pc", "price")),
        ("quantity", ("qty", "quantity")),
        ("load_speed_rating", ("l.i", "li/sr", "li&sr", "load")),
        ("item_code", ("code",)),
        ("item_no", ("no.", "no")),
        ("brand", ("brand",)),
        ("pattern", ("pattern",)),
        ("size", ("size",)),
    ]

    @staticmethod
    def _clean_value(value: Any) -> Any:
        if value is None:
            return None

        if isinstance(value, (str, int, float, bool)):
            return value

        try:
            return value.isoformat()
        except Exception:
            return str(value)

    @classmethod
    def _iter_rows(
        cls,
        worksheet,
        value_worksheet,
    ):
        """
        Yield cleaned cell values per row.

        Formula cells (e.g. =H7*I7) use the value cached by Excel
        from value_worksheet. The formula text is kept only when no
        cached value exists, so parsers can still compute it.
        """

        for row, value_row in zip(
            worksheet.iter_rows(values_only=True),
            value_worksheet.iter_rows(values_only=True),
        ):
            yield [
                cls._clean_value(
                    cached
                    if cached is not None
                    else value
                )
                for value, cached in zip(row, value_row)
            ]

    @staticmethod
    def _row_to_text(values: list[Any]) -> str:
        return " | ".join(
            "" if value is None else str(value).strip()
            for value in values
        )

    @staticmethod
    def _detect_excel_format(
        header_values: list[Any],
    ) -> str:

        headers = [
            ""
            if value is None
            else str(value).strip().lower()
            for value in header_values
        ]

        # ---------------------------------------------------------
        # Compact format
        # ---------------------------------------------------------
        #
        # Actual workbook:
        #
        # SIZE AND PATTERN
        # Quantity
        # [blank]
        # Unit Price FOB Laem Chabang
        # Amount
        # [blank]
        # Qty in 40HQ
        # [blank]
        # [blank]
        #
        if (
            any(
                "size and pattern" in value
                for value in headers
            )
            and any(
                value == "quantity"
                for value in headers
            )
            and any(
                "unit price" in value
                for value in headers
            )
            and any(
                "amount" in value
                for value in headers
            )
            and any(
                "qty in 40hq" in value
                for value in headers
            )
        ):
            return "compact"

        # ---------------------------------------------------------
        # Existing structured format
        # ---------------------------------------------------------

        if (
            any(
                value in {"no.", "no"}
                for value in headers
            )
            and any(
                "specification" in value
                for value in headers
            )
        ):
            return "structured"

        # ---------------------------------------------------------
        # Header-mapped format
        # ---------------------------------------------------------
        #
        # No. | CODE | Brand | Pattern | Size | L.I&S.R |
        # Order Qty | QTY/40HQ | FOB (US$/PC) | Total Amount
        #
        if (
            any("pattern" in value for value in headers)
            and any("size" in value for value in headers)
            and any("brand" in value for value in headers)
            and any(
                "qty" in value or "quantity" in value
                for value in headers
            )
        ):
            return "mapped"

        return "unknown"


    @classmethod
    def _find_table_format(
        cls,
        rows: list[list[Any]],
    ) -> tuple[int | None, str]:
        """
        Find the first recognizable table header.

        Returns:
            (header_row_number, format_name)
        """

        for row_number, values in enumerate(
            rows,
            start=1,
        ):
            if not any(value is not None for value in values):
                continue

            excel_format = cls._detect_excel_format(values)

            if excel_format != "unknown":
                return row_number, excel_format

            # Two-row header, e.g.:
            #
            # COMMODITY |   |   |         |      |         | Order Qty | QTY/40HQ
            # No.       | CODE | Brand | Pattern | Size | L.I&S.R |           |

            if row_number >= 2:

                merged = cls._merge_header_rows(
                    rows[row_number - 2],
                    values,
                )

                if cls._detect_excel_format(merged) == "mapped":
                    return row_number, "mapped"

        return None, "unknown"

    @staticmethod
    def _merge_header_rows(
        upper: list[Any],
        lower: list[Any],
    ) -> list[Any]:

        merged = []

        for index in range(max(len(upper), len(lower))):

            parts = [
                str(row[index]).strip()
                for row in (upper, lower)
                if index < len(row)
                and row[index] is not None
                and str(row[index]).strip()
            ]

            merged.append(" ".join(parts) or None)

        return merged

    @classmethod
    def _map_columns(
        cls,
        header_values: list[Any],
    ) -> dict[str, int]:

        columns: dict[str, int] = {}

        for index, value in enumerate(header_values):

            if value is None:
                continue

            header = " ".join(str(value).lower().split())

            for field, keywords in cls.COLUMN_KEYWORDS:

                if field in columns:
                    continue

                if any(
                    header == keyword
                    if keyword == "no"
                    else keyword in header
                    for keyword in keywords
                ):
                    columns[field] = index
                    break

        return columns

    @classmethod
    def _build_ocr_block(
        cls,
        values: list[Any],
        row_number: int,
        excel_format: str,
        columns: dict[str, int] | None = None,
    ) -> OCRBlock:

        text = cls._row_to_text(values)

        return OCRBlock(
            text=text,
            page=1,
            x1=0.0,
            y1=float(row_number),
            x2=float(max(len(values), 1)),
            y2=float(row_number + 1),
            confidence=1.0,
            source="excel",
            metadata={
                "row_number": row_number,
                "values": values,
                "excel_format": excel_format,
                "columns": columns or {},
            },
        )

    def _extract_detected_tables(
        self,
        rows: list[list[Any]],
    ) -> list[Table]:
        """
        Known sheet formats compete with the layout-independent
        header mapping; the table with the most product data wins.
        Without any product table, the sheet is returned as found.
        """

        candidates = self._extract_format_tables(rows)

        generic_rows = GenericTableExtractor().rows_from_grid(
            rows,
            carry_over=False,
        )

        if generic_rows:
            candidates.append(
                rows_to_table(
                    generic_rows,
                    "detected_table_1",
                    None,
                    "grid",
                    self.HEADERS,
                )
            )

        if candidates:

            # Known sheet formats come first and win ties.
            best = max(
                candidates,
                key=lambda table: table_score(table)[:2],
            )

            if table_score(best)[1] > 0:
                return [best]

        table = raw_table(rows, "raw_table_1", None)

        return [table] if table is not None else []

    def _extract_format_tables(
        self,
        rows: list[list[Any]],
    ) -> list[Table]:

        header_row, excel_format = self._find_table_format(
            rows
        )

        if header_row is None:
            return []

        columns = None

        if excel_format == "mapped":

            header_values = rows[header_row - 1]

            if header_row >= 2:
                header_values = self._merge_header_rows(
                    rows[header_row - 2],
                    header_values,
                )

            columns = self._map_columns(header_values)

        blocks: list[OCRBlock] = []

        for row_number, values in enumerate(
            rows,
            start=1,
        ):
            if row_number <= header_row:
                continue

            if not any(value is not None for value in values):
                continue

            block = self._build_ocr_block(
                values=values,
                row_number=row_number,
                excel_format=excel_format,
                columns=columns,
            )

            blocks.append(block)

        parsed_rows = self.table_service.parse_rows(
            blocks
        )

        if not parsed_rows:
            return []

        table_rows = [
            [
                table_value(row.cells, header)
                for header in self.HEADERS
            ]
            for row in parsed_rows
        ]

        return [
            Table(
                table_id="detected_table_1",
                page=None,
                headers=self.HEADERS,
                rows=table_rows,
                extraction_method=f"excel_{excel_format}",
            )
        ]

    def _load_xlsx(
        self,
        file_path: str,
    ) -> list[tuple[str, int, int, list[list[Any]]]]:

        workbook = load_workbook(
            file_path,
            data_only=False,
            read_only=False,
        )

        # Same workbook with the values Excel cached for formulas.
        value_workbook = load_workbook(
            file_path,
            data_only=True,
            read_only=False,
        )

        loaded = [
            (
                worksheet.title,
                worksheet.max_row or 0,
                worksheet.max_column or 0,
                evaluate_formulas(
                    list(
                        self._iter_rows(
                            worksheet,
                            value_worksheet,
                        )
                    )
                ),
            )
            for worksheet, value_worksheet in zip(
                workbook.worksheets,
                value_workbook.worksheets,
            )
        ]

        workbook.close()
        value_workbook.close()

        return loaded

    def _load_xls(
        self,
        file_path: str,
    ) -> list[tuple[str, int, int, list[list[Any]]]]:
        """
        Legacy .xls (BIFF) workbook.

        xlrd returns cached formula values, "" for empty cells and
        floats for every number (850 -> 850.0), so values are
        converted to match the openpyxl output.
        """

        workbook = xlrd.open_workbook(file_path)

        loaded = []

        for worksheet in workbook.sheets():

            all_rows = []

            for row_index in range(worksheet.nrows):

                values = []

                for cell in worksheet.row(row_index):

                    value = cell.value

                    if cell.ctype in (
                        xlrd.XL_CELL_EMPTY,
                        xlrd.XL_CELL_BLANK,
                    ) or value == "":
                        value = None

                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(
                            value,
                            workbook.datemode,
                        )

                    elif (
                        isinstance(value, float)
                        and value.is_integer()
                    ):
                        value = int(value)

                    values.append(self._clean_value(value))

                all_rows.append(values)

            loaded.append(
                (
                    worksheet.name,
                    worksheet.nrows,
                    worksheet.ncols,
                    all_rows,
                )
            )

        workbook.release_resources()

        return loaded

    def process(
        self,
        file_path: str,
        document_id: str,
        file_name: str,
        mime_type: str,
    ) -> DocumentResult:

        sheets: list[Sheet] = []
        all_text: list[str] = []

        if mime_type == "application/vnd.ms-excel":
            loaded = self._load_xls(file_path)
            engine = "xlrd + shared_table_detector"
        else:
            loaded = self._load_xlsx(file_path)
            engine = "openpyxl + shared_table_detector"

        for name, max_row, max_column, all_rows in loaded:

            rows: list[list[Any]] = []

            for cleaned in all_rows:

                if any(
                    value is not None
                    for value in cleaned
                ):
                    rows.append(cleaned)

                    all_text.append(
                        " | ".join(
                            ""
                            if value is None
                            else str(value)
                            for value in cleaned
                        )
                    )

            tables = self._extract_detected_tables(
                all_rows
            )

            sheets.append(
                Sheet(
                    name=name,
                    max_row=max_row,
                    max_column=max_column,
                    rows=rows,
                    tables=tables,
                )
            )

        # Raw tables only when no sheet has a product table.
        if any(
            table.extraction_method != "raw"
            for sheet in sheets
            for table in sheet.tables
        ):
            for sheet in sheets:
                sheet.tables = [
                    table
                    for table in sheet.tables
                    if table.extraction_method != "raw"
                ]

        return DocumentResult(
            document_id=document_id,
            file_name=file_name,
            mime_type=mime_type,
            document_type="excel",
            text="\n".join(all_text),
            sheets=sheets,
            metadata={
                "engine": engine,
                "table_detection": "shared",
            },
        )
